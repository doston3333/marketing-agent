#!/usr/bin/env python3
"""Command line for the AI Station tender agent. It opens the tender sites in data/tender_sites.json (signing in
where credentials are set), remembers every lot it has seen (data/tenders.json) and sends Claude's review of
the new ones to Telegram. Separate from the marketing agent: own bot, own data files.

Setup + status
  check                       secrets, bot, browser, and which sites have / need login credentials
  sites                       the site registry with the result of each site's last scan
  login SITE                  try signing in to one site and report the result

Collecting
  scan [--site ID ...]        open every enabled site, extract its lots, record the new ones (prints a summary)
  new [--min-score N] [--limit N] [--all]
                              lots not reviewed yet: keyword matches in full, the rest as one-line titles
  open KEY|URL [--site ID] [--max N] [--click LABEL ...]
                              open one tender page (logged in when the site has credentials): text, the text
                              behind the site's detail tabs (detail_tabs in the registry) and document links
  doc URL|KEY [--click LABEL [--nth N]] [--site ID] [--max N]
                              download a tender document (pdf/docx/xlsx) and print its text; with --click, URL is
                              the tender page and LABEL its download button

Reviewing + delivering
  profile                     what AI Station bids on, keyword lists, and the team's feedback so far
  review FILE [--keep-rest]   store Claude's verdicts (JSON list); the other lots the last `new` showed are marked
                              seen unless --keep-rest (lots `new` did not show stay new for the next run)
  digest FILE [--dry-run]     send the Telegram digest for the reviewed lots in FILE (--dry-run prints it)
  say HTML                    send a free-form message to the tender recipients
  inbox                       the team's new replies to tender messages (routed by the marketing agent from the
                              shared bot; with TENDER_BOT_TOKEN, read from that bot instead)
  feedback add TEXT | feedback list
  save                        commit this agent's data files to the repo's default branch (GitHub API)
"""
import argparse, datetime, json, os, re, subprocess, sys, time

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "lib"))
import agent_lib as A  # noqa: E402
import tenders as T  # noqa: E402

NO_BROWSER_MODES = {"worldbank", "api"}
DOC_EXT = re.compile(r"\.(pdf|docx?|xlsx?|zip|rar|rtf|odt)(\?|$)", re.I)


def _browser():
    import browser as B
    return B.Browser()


def _enabled(only=None):
    out = [s for s in T.sites() if s.get("enabled", True)]
    if only:
        out = [T.site_by_id(x) for x in only]
    return out


def cmd_check(a):
    import browser as B
    st = T.load_state()
    print("TENDER_BOT_TOKEN:", "set" if T.bot_token() else "not set (digest falls back to the marketing bot, send-only)")
    print("TELEGRAM_BOT_TOKEN:", "set" if A.secrets().get("telegram_bot_token") else "MISSING")
    try:
        me = T.tg("getMe", {})["result"]
        print("bot:", "@" + me.get("username", "?"))
    except Exception as ex:
        print("bot: ERROR", ex)
    print("recipients:", T.recipients(st) or "NONE (press Start in the tender bot, or set TENDER_CHAT_IDS)")
    try:
        import playwright  # noqa: F401
        print("playwright: installed;", json.dumps(B.setup()))
    except ImportError:
        print("playwright: MISSING (pip install -r requirements-tender.txt)")
    print("items remembered:", len(st["items"]), "| last run:", (st["runs"][-1:] or [{}])[0].get("date"))
    for s in T.sites():
        u, p = B.env_names(s)
        have = all(os.environ.get(x) for x in (u, p))
        login = "login: credentials set" if have and s.get("login") else \
            (f"login: optional, set {u}/{p}" if s.get("login") else "public")
        print(f"  {'on ' if s.get('enabled', True) else 'off'} {s['id']:<18} {login}")


def cmd_sites(a):
    st = T.load_state()
    for s in T.sites():
        ss = st["site_status"].get(s["id"], {})
        when = time.strftime("%m-%d %H:%M", time.gmtime(ss["at"] + 5 * 3600)) if ss.get("at") else "never"
        print(f"{'on ' if s.get('enabled', True) else 'off'} {s['id']:<18} {s['name']:<30} last scan {when}: "
              f"{ss.get('found', '-')} lots, {ss.get('new', '-')} new {('| ' + '; '.join(ss['notes'])) if ss.get('notes') else ''}")


def cmd_scan(a):
    st = T.load_state()
    targets = _enabled(a.site)
    problems, total_new, ok = [], 0, 0
    br = None
    try:
        modes = {(p.get("mode") if isinstance(p, dict) else None) or s.get("mode", "links")
                 for s in targets for p in s["pages"]}
        if modes - NO_BROWSER_MODES:
            br = _browser()
        for s in targets:
            items, notes = T.collect_site(s, br)
            new = T.merge(st, s, items, notes)
            total_new += len(new)
            hits = sum(1 for k in new if st["items"][k]["score"] > 0)
            closed = st["site_status"][s["id"]]["closed"]
            ok += bool(items) or not notes
            if notes and not items:
                problems.append(f"{s['name']}: {notes[0]}")
            print(f"{s['id']:<18} {len(items):>5} lots {closed:>4} closed {len(new):>5} new {hits:>4} keyword hits  "
                  f"{('! ' + '; '.join(notes)) if notes else ''}", flush=True)
            T.save_state(st)  # after every site, so a crash later keeps what was collected
    finally:
        if br:
            br.close()
    st["runs"].append({"date": A.tashkent_now().strftime("%Y-%m-%d %H:%M"), "new": total_new,
                       "sites": len(targets), "sites_ok": ok,
                       "problems": problems, "reviewed": []})
    T.save_state(st)
    print(json.dumps({"new_total": total_new, "open_unreviewed": len(T.open_items(st)),
                      "problems": problems}, ensure_ascii=False))


def cmd_new(a):
    st = T.load_state()
    items = T.open_items(st)
    if not a.all:
        items = [i for i in items if not (i.get("deadline") and i["deadline"] < datetime.date.today().isoformat())]
    matches = [i for i in items if i.get("score", 0) >= a.min_score]
    others = [i for i in items if i.get("score", 0) < a.min_score]
    out = {"open_unreviewed": len(items), "keyword_matches": len(matches),
           "matches": [{"key": i["key"], "score": i["score"], "hits": i["hits"], "title": i["title"],
                        "deadline": i.get("deadline"), "url": i["url"], **(i.get("detail") or {})}
                       for i in matches[:a.limit]],
           "others": [f"{i['key']} | {i['title'][:110]} | {i.get('deadline') or '-'}" for i in others[:a.others]],
           "others_not_shown": max(0, len(others) - a.others)}
    # review marks only what was shown here as seen; the rest stays new for the next run
    st["shown"] = [i["key"] for i in matches[:a.limit] + others[:a.others]]
    T.save_state(st)
    print(json.dumps(out, ensure_ascii=False, indent=1))


def _resolve(target, site_id=None):
    st = T.load_state()
    if target in st["items"]:
        it = st["items"][target]
        return T.site_by_id(it["site"]), it["url"]
    if site_id:
        return T.site_by_id(site_id), target
    host = re.sub(r"^https?://(www\.)?", "", target).split("/")[0]
    for s in T.sites():
        urls = [p if isinstance(p, str) else p["url"] for p in s["pages"]] + [s.get("url", "")]
        if any(host and host in u for u in urls):
            return s, target
    return {"id": "web", "name": "web"}, target


def cmd_open(a):
    site, url = _resolve(a.target, a.site)
    with _browser() as br:
        res = br.fetch(site, url, scroll=1, tabs=(a.click or []) + site.get("detail_tabs", []))
        login = br.logins.get(site["id"])
    os.makedirs(T.PAGES_DIR, exist_ok=True)
    path = os.path.join(T.PAGES_DIR, "open-" + re.sub(r"\W+", "_", url)[-80:] + ".txt")
    with open(path, "w") as f:
        f.write(res["text"])
    docs, seen = [], set()
    skip = re.compile(site["doc_exclude"]) if site.get("doc_exclude") else None
    for l in res["links"]:
        if skip and skip.search(l["href"]):
            continue
        if DOC_EXT.search(l["href"]) or re.search(r"download|yuklab|скачать|attachment|document", l["text"], re.I):
            if l["href"] not in seen:
                seen.add(l["href"])
                docs.append({"text": l["text"][:120], "url": l["href"]})
    text = re.sub(r"\n{3,}", "\n\n", res["text"])
    print(json.dumps({"url": res["url"], "title": res["title"], "site": site["id"], "login": login,
                      "blocked": res["blocked"], "error": res["error"], "documents": docs[:30],
                      "download_buttons": sorted(set(res.get("download_buttons") or [])),
                      "full_text_saved_to": path, "text_chars": len(text)}, ensure_ascii=False, indent=1))
    print("\n----- PAGE TEXT -----")
    print(text[:a.max] + ("\n[... cut, read the saved file for the rest]" if len(text) > a.max else ""))


def doc_text(path):
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True)
        if r.returncode == 0:
            return r.stdout
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    if ext == "docx":
        import docx
        d = docx.Document(path)
        rows = [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
        return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(rows)
    if ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return "\n".join(" | ".join("" if v is None else str(v) for v in row)
                         for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    return f"(no text extractor for .{ext}; file saved at {path})"


def cmd_doc(a):
    site, url = _resolve(a.url, a.site)
    d = os.path.join(T.PAGES_DIR, "docs")
    os.makedirs(d, exist_ok=True)
    name = re.sub(r"[^\w.\-]+", "_", url.split("?")[0].rstrip("/").split("/")[-1])[-80:] or "document"
    with _browser() as br:
        if a.click:  # the file sits behind a button on the tender page (url is that page)
            path = br.download_by_click(site, url, a.click, d, nth=a.nth)
            name = os.path.basename(path)
        else:
            path = br.download(site, url, os.path.join(d, name))
    if "." not in name:  # no extension in the URL: sniff the bytes
        head = open(path, "rb").read(4)
        ext = "pdf" if head.startswith(b"%PDF") else ("docx" if head.startswith(b"PK") else "bin")
        os.rename(path, path + "." + ext)
        path += "." + ext
    text = doc_text(path)
    print(f"saved: {path} ({os.path.getsize(path) // 1024} KB, {len(text)} chars of text)\n")
    if len(text.strip()) < 200 and path.lower().endswith(".pdf"):
        print("This looks like a scanned PDF (no text layer): open the saved file with the Read tool "
              "(pages \"1-5\" and so on) to read it as images.")
    print(text[:a.max] + ("\n[... cut]" if len(text) > a.max else ""))


def cmd_profile(a):
    st = T.load_state()
    prof = A.drive_get(T.PROFILE_FILE)
    print(prof.decode() if prof else "(no data/tender_profile.md)")
    kw = T.keywords()
    print("\n## Keyword lists (data/tender_keywords.json)")
    for k in ("strong", "medium", "negative"):
        print(f"{k}: {', '.join(kw.get(k, []))}")
    print("\n## Team feedback (newest last; it overrides the profile)")
    for f in st["feedback"][-30:]:
        when = time.strftime("%Y-%m-%d", time.gmtime(f["ts"]))
        ctx = f" (reply to: {f['reply_to'][:200]!r})" if f.get("reply_to") else ""
        print(f"- {when} {f.get('from', '')}: {f['text']}{ctx}")
    recent = [it for it in st["items"].values() if it.get("verdict") in ("bid", "watch")][-15:]
    if recent:
        print("\n## Recently recommended (do not repeat unless something changed)")
        for it in recent:
            print(f"- [{it['verdict']}] {it['title'][:120]} (deadline {it.get('deadline')})")


def cmd_review(a):
    st = T.load_state()
    reviews = json.load(open(a.file))
    if isinstance(reviews, dict):
        reviews = reviews.get("reviews", [])
    now, done = time.time(), []
    for r in reviews:
        it = st["items"].get(r.get("key"))
        if not it and r.get("url") and r.get("title"):  # found outside the scan (e.g. web search on a blocked site)
            r["key"] = r.get("key") or "web:" + T._hid(r["url"])
            it = st["items"].setdefault(r["key"], {
                "site": r.get("site", "web"), "title": r["title"][:200], "url": r["url"],
                "deadline": T.parse_date(r.get("deadline")), "first_seen": now, "last_seen": now,
                "score": 0, "hits": [], "status": "new"})
        if not it:
            print("unknown key (give url + title for lots found outside the scan), skipped:", r.get("key"))
            continue
        it.update({"status": "reviewed", "verdict": r.get("verdict", "skip"), "fit": r.get("fit"),
                   "reviewed_at": now})
        if r.get("deadline"):
            it["deadline"] = T.parse_date(r["deadline"]) or it.get("deadline")
        if r.get("verdict") in T.VERDICTS:
            it["review"] = {k: r[k] for k in ("title_uz", "summary", "why", "next_step", "amount", "buyer") if r.get(k)}
        done.append(r["key"])
    rest = 0
    if not a.keep_rest:
        for k in st.get("shown", []):
            it = st["items"].get(k)
            if it and it.get("status") == "new":
                it["status"], rest = "seen", rest + 1
        st["shown"] = []
    if st["runs"]:
        st["runs"][-1]["reviewed"] = done
    T.save_state(st)
    counts = {}
    for r in reviews:
        counts[r.get("verdict", "skip")] = counts.get(r.get("verdict", "skip"), 0) + 1
    print(json.dumps({"stored": len(done), "verdicts": counts, "other_shown_marked_seen": rest,
                      "still_new": sum(1 for it in st["items"].values() if it.get("status") == "new")}))


def cmd_digest(a):
    st = T.load_state()
    reviews = json.load(open(a.file))
    if isinstance(reviews, dict):
        reviews = reviews.get("reviews", [])
    run = st["runs"][-1] if st["runs"] else {"new": 0, "sites": 0, "sites_ok": 0, "problems": []}
    text = T.digest_html(st, reviews, run)
    if a.dry_run:
        print(text)
        return
    mids = T.send_html(st, text)
    T.save_state(st)  # keeps the sent message ids, so replies to them are routed back here
    print(f"sent {len(mids)} message(s) to {T.recipients(st)}")


def cmd_save(a):
    import ghstore
    try:
        print(ghstore.save(f"tender agent: state {A.tashkent_now():%Y-%m-%d %H:%M}", only=T.OWN_FILES))
    except Exception as ex:
        print("GitHub save failed:", ex)
        sys.exit(1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for c in ("check", "sites", "profile", "inbox", "save"):
        sub.add_parser(c)
    p = sub.add_parser("login"); p.add_argument("site")
    p = sub.add_parser("scan"); p.add_argument("--site", nargs="*")
    p = sub.add_parser("new"); p.add_argument("--min-score", type=int, default=1); p.add_argument("--limit", type=int, default=150)
    p.add_argument("--others", type=int, default=400); p.add_argument("--all", action="store_true")
    p = sub.add_parser("open"); p.add_argument("target"); p.add_argument("--site"); p.add_argument("--max", type=int, default=12000)
    p.add_argument("--click", action="append", help="tab or section label to click and read too (repeatable)")
    p = sub.add_parser("doc"); p.add_argument("url"); p.add_argument("--site"); p.add_argument("--max", type=int, default=15000)
    p.add_argument("--click", help="label of the download button on the page at URL (see open's download_buttons)")
    p.add_argument("--nth", type=int, default=0, help="which of several buttons with that label (0 = first)")
    p = sub.add_parser("review"); p.add_argument("file"); p.add_argument("--keep-rest", action="store_true")
    p = sub.add_parser("digest"); p.add_argument("file"); p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("say"); p.add_argument("html")
    p = sub.add_parser("feedback"); p.add_argument("op", choices=["add", "list"]); p.add_argument("text", nargs="?")
    a = ap.parse_args()

    if a.cmd == "login":
        s = T.site_by_id(a.site)
        with _browser() as br:
            print(s["id"], br.login(s, force=True))
        return
    if a.cmd == "inbox":
        st = T.load_state()
        res = T.inbox(st)
        T.save_state(st)
        print(json.dumps(res, ensure_ascii=False, indent=1))
        return
    if a.cmd == "say":
        st = T.load_state()
        n = len(T.send_html(st, a.html))
        T.save_state(st)
        print("sent", n, "message(s)")
        return
    if a.cmd == "feedback":
        st = T.load_state()
        if a.op == "add":
            st["feedback"].append({"ts": time.time(), "from": "agent", "text": a.text})
            T.save_state(st)
        for f in st["feedback"][-30:]:
            print("-", f.get("from"), ":", f["text"])
        return
    {"check": cmd_check, "sites": cmd_sites, "scan": cmd_scan, "new": cmd_new, "open": cmd_open, "doc": cmd_doc,
     "profile": cmd_profile, "review": cmd_review, "digest": cmd_digest, "save": cmd_save}[a.cmd](a)


if __name__ == "__main__":
    main()
